"""Create the editable website Word from mapa-textos.json; never touch informe/.

Requires python-docx. First run: node scripts/export-copy.mjs (from web/).
Usage: python scripts/export-copy-word.py
"""

import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path(__file__).resolve().parents[1]
MAP = BASE / "textos/mapa-textos.json"
OUT = BASE / "textos/textos-pagina.docx"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def main():
    if OUT.exists():
        raise SystemExit(f"Ya existe {OUT}. No sobrescribir una posible edición de José.")
    data = json.loads(MAP.read_text())
    for name, metadata in data["files"].items():
        digest = hashlib.sha256((BASE / name).read_bytes()).hexdigest()
        if digest != metadata["sha256"]:
            raise SystemExit(f"La página cambió desde la extracción: {name}")

    doc = Document()
    doc.core_properties.title = "Textos de la página — Una sola gravedad, tres escalas"
    doc.core_properties.subject = "Edición de textos de la web; documento independiente del informe AST101"
    doc.core_properties.author = "José Santiago Garrillo Fernández"
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    caption = doc.styles["Caption"]
    caption.font.size = Pt(8)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string("666666")
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_before = Pt(8)
    caption.paragraph_format.space_after = Pt(3)

    doc.add_heading("Textos de la página", 0)
    doc.add_paragraph("Una sola gravedad, tres escalas", "Subtitle")
    doc.add_heading("Cómo editar", 1)
    for text in [
        "Este Word contiene los textos de la web. Es independiente del informe: no modifica ni reemplaza el documento de AST101.",
        "Edita directamente los textos bajo cada identificador (por ejemplo, SOL-009). Puedes reescribir frases y agregar o quitar párrafos dentro del mismo bloque. No necesitas editar HTML ni código.",
        "Conserva los identificadores y las marcas amarillas entre ⟦ y ⟧: son valores que calcula o elige el programa, no palabras para reemplazar por números fijos. Puedes mover una marca dentro de su frase, pero no cambiar su nombre.",
        "También están los controles, las tablas, los ejes, las leyendas, los mensajes de carga y los textos accesibles. Algunos rótulos se repiten porque aparecen en lugares distintos. Los nombres de modelos compartidos se definen en Universo y también se usan en Supernovas.",
        "Los bloques de «Texto alternativo» y los que siguen inmediatamente a un texto con marcas contienen variantes que la página muestra según sus datos o controles. No es necesario modificar las unidades, las fórmulas ni los nombres de planetas si solo buscas cambiar la redacción.",
        "Cuando termines, guarda el archivo como .docx y devuélvemelo o indícame su ruta. Compararé tus cambios con esta copia, los aplicaré a la página, comprobaré que siga funcionando y la resubiré. El informe queda fuera de ese proceso.",
    ]:
        doc.add_paragraph(text)
    count = sum(len(g["entries"]) for g in data["groups"])
    doc.add_paragraph(f"Inventario: {count} bloques de texto, organizados en {len(data['groups'])} secciones.")

    for index, group in enumerate(data["groups"]):
        # Major tabs start on a fresh page; the short ancillary sections flow.
        if index < 5:
            doc.add_page_break()
        doc.add_heading(group["title"], 1)
        for entry in group["entries"]:
            doc.add_paragraph(f"{entry['id']} · {entry['kind']}", "Caption")
            p = doc.add_paragraph()
            for chunk in re.split(r"(⟦[^⟧]+⟧)", entry["text"]):
                if not chunk:
                    continue
                run = p.add_run(chunk)
                if chunk.startswith("⟦"):
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            # Native rich-text content controls retain IDs even if paragraphs
            # are added, moved or restyled inside Word. They contain no code.
            control = OxmlElement("w:sdt")
            properties = OxmlElement("w:sdtPr")
            tag = OxmlElement("w:tag")
            tag.set(qn("w:val"), entry["id"])
            alias = OxmlElement("w:alias")
            alias.set(qn("w:val"), f"{entry['id']} — {entry['kind']}")
            properties.extend([tag, alias])
            content = OxmlElement("w:sdtContent")
            parent = p._p.getparent()
            parent.replace(p._p, control)
            content.append(p._p)
            control.extend([properties, content])

    doc.save(OUT)
    # Verify every editable block and every placeholder survived the DOCX.
    with ZipFile(OUT) as archive:
        assert archive.testzip() is None
        from lxml import etree
        root = etree.fromstring(archive.read("word/document.xml"))
        controls = root.findall(f".//{{{W}}}sdt")
        assert len(controls) == count
        actual = {}
        for control in controls:
            key = control.find(f"{{{W}}}sdtPr/{{{W}}}tag").get(f"{{{W}}}val")
            actual[key] = "".join(control.itertext())
        for group in data["groups"]:
            for entry in group["entries"]:
                assert actual[entry["id"]] == entry["text"], entry["id"]
    print(f"Word creado: {OUT} ({count} bloques verificados)")


if __name__ == "__main__":
    main()
